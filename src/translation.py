# Import necessary libraries
from googletrans import Translator
from PyPDF2 import PdfReader
import cv2
import pytesseract
from fpdf import FPDF
from googletrans import Translator
from docx import Document
import textwrap
import time 

# Define a class for handling translation and file operations
class TranslationFile:
    # Mapping of language names to language codes for translation
    LANGUAGE_MAPPING = {
        'telugu': 'te',
        'english': 'en',
        'assamese': 'as',
        # ... (add other languages if needed)
    }

    # Function to extract text from an image using OCR
    def extract_text_from_image(self, scanned_image):
        # Convert the scanned image to grayscale
        gray_image = cv2.cvtColor(scanned_image, cv2.COLOR_RGB2GRAY)
        # Use pytesseract to extract text from the grayscale image
        text = pytesseract.image_to_string(gray_image)
        print("Extracted Text from Image:")
        print(text)
        return text

    # Function to extract text from a PDF file
    def extract_text_from_pdf(self, pdf_path):
        # Open the PDF file in binary mode
        with open(pdf_path, 'rb') as file:
            # Create a PdfReader object to read the PDF file
            pdf_reader = PdfReader(file)
            text = ""
            # Iterate through each page of the PDF and extract text
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                text += page.extract_text()
            print("Extracted Text from PDF:")
            return text

    # Function to extract text from a plain text file
    def extract_text_from_text_file(self, text_file_path):
        # Open the text file and read its contents
        with open(text_file_path, 'r') as file:
            text = file.read()
            print("Text from Text File:")
            print(text)
            return text

    # Function to extract text from a Word document
    def extract_text_from_word_document(self, docx_path):
        # Use the python-docx library to read text from a Word document
        doc = Document(docx_path)
        text = ""
        # Iterate through each paragraph in the document and extract text
        for paragraph in doc.paragraphs:
            text += paragraph.text + '\n'
        print("Text from Word Document:")
        print(text)
        return text

    # Function to translate text to a specified target language
    def translate_text(self, text, target_language):
        # Create a Translator object from the googletrans library
        translator = Translator()

        # Handle empty text
        if not text:
            print("Error: Empty text provided for translation.")
            return ""
       
        try:
            # Use the translator to translate the text to the target language
            translation = translator.translate(text, dest=target_language)
            translated_text = translation.text
            print("Translated Text:")
            return translated_text
        except Exception as e:
            print(f"Error during translation: {e}")
            return ""

    # Function to convert text to PDF format
    def text_to_pdf(self, text, filename="summarized_text.pdf"):
        # Set up parameters for creating a PDF using FPDF
        a4_width_mm = 210
        pt_to_mm = 0.35
        fontsize_pt = 10
        fontsize_mm = fontsize_pt * pt_to_mm
        margin_bottom_mm = 10
        character_width_mm = 7 * pt_to_mm
        width_text = a4_width_mm / character_width_mm

        # Create an FPDF object with A4 dimensions
        pdf = FPDF(orientation='P', unit='mm', format='A4')
        pdf.set_auto_page_break(True, margin=margin_bottom_mm)
        pdf.add_page()
        pdf.set_font(family='Courier', size=fontsize_pt)
        splitted = text.split('\n')

        # Wrap text and add it to the PDF
        for line in splitted:
            lines = textwrap.wrap(line, width_text)

            if len(lines) == 0:
                pdf.ln()

            for wrap in lines:
                pdf.cell(0, fontsize_mm, wrap.encode('latin-1', 'replace').decode('latin-1'), ln=1)

        # Output the PDF to a file
        pdf.output(filename, 'F')
        return filename 

# Function to chunk the text into smaller pieces and translate each chunk
def chunk_and_translate(text, chunk_size, target_language):
    chunks = [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]

    translated_chunks = []
    for chunk in chunks:
        translated_chunk = translation_text.translate_text(chunk, target_language)
        translated_chunks.append(translated_chunk)

    return ''.join(translated_chunks)

# Create an instance of the TranslationFile class
translation_text = TranslationFile()

# Measure the execution time
start_time = time.time()

# Specify the file path and extract text from the PDF file
file = "CBN Sir Bail Judgement Copy.pdf"
text = translation_text.extract_text_from_pdf(file)

# Specify the target language and chunk size for translation
target_language = 'te'
chunk_size = 4500

# Translate the text in chunks
translated_text = chunk_and_translate(text, chunk_size, target_language)

# Measure the total execution time
end_time = time.time()
elapsed_time = end_time - start_time
print(f"Total execution time: {elapsed_time} seconds")

# Display the total extracted text
print("\nTotal Extracted Text:")
print(text)


# Display the total translated text
print("\nTotal Translated Text:")
print(translated_text)

